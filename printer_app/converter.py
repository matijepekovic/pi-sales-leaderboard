"""Render a newly-created, values-only workbook with a private LibreOffice profile."""
from __future__ import annotations

import math
import os
import signal
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from .config import Config
from .parser import Group


class ConversionError(ValueError):
    pass


def page_count(path: Path) -> int:
    try:
        with path.open('rb') as stream:
            reader = PdfReader(stream, strict=False)
            if reader.is_encrypted:
                raise ConversionError('ENCRYPTED PDF IS NOT PRINTABLE')
            count = len(reader.pages)
            if count < 1:
                raise ConversionError('NO PRINTABLE PDF PAGES FOUND')
            return count
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError('INVALID PDF: ' + type(exc).__name__) from exc


def write_workbook(group: Group, target: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = 'Report'
    for row in [group.headers, *group.rows]:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            # Even a cached formula result that begins '=' remains literal text.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = 'mm/dd/yyyy'
            cell.font = Font(name='Liberation Sans', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    widths = []
    for c in range(1, sheet.max_column + 1):
        longest = max(len(str(sheet.cell(r, c).value or '')) for r in range(1, sheet.max_row + 1))
        width = min(36, max(10, min(longest + 2, 28)))
        widths.append(width)
        sheet.column_dimensions[get_column_letter(c)].width = width
        header = sheet.cell(1, c)
        header.font = Font(name='Liberation Sans', size=10, bold=True, color='FFFFFF')
        header.fill = PatternFill('solid', fgColor='26384B')
    # Explicit row heights prevent LibreOffice clipping wrapped cells. Do not cap:
    # a very long row must create excess pages and an error, not hidden data.
    for r in range(1, sheet.max_row + 1):
        lines = max(sum(max(1, math.ceil(len(line) / max(1, widths[c - 1] - 2)))
                        for line in str(sheet.cell(r, c).value or '').split('\n'))
                    for c in range(1, sheet.max_column + 1))
        if lines * 15 + 5 > 400:
            raise ConversionError('CELL TEXT IS TOO LONG FOR A READABLE REPORT')
        sheet.row_dimensions[r].height = max(22, lines * 15 + 5)
    sheet.print_area = f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    sheet.print_title_rows = '1:1'
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0  # Never conceal excess pages with unlimited shrinking.
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.3, bottom=.3, header=.1, footer=.1)
    sheet.print_options.horizontalCentered = True
    book.save(target)
    book.close()


def convert(workbook: Path, directory: Path, cfg: Config) -> Path:
    target = directory / (workbook.stem + '.pdf')
    target.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix='lo-', dir=directory) as temp:
        home = Path(temp)
        profile = home / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user' / 'registrymodifications.xcu').write_text('''<?xml version="1.0"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
</oor:items>''')
        command = [cfg.libreoffice, f'-env:UserInstallation={profile.as_uri()}', '--headless',
                   '--nologo', '--nodefault', '--norestore', '--convert-to', 'pdf:calc_pdf_Export',
                   '--outdir', str(directory), str(workbook)]
        # Credentials and the user's Office profile are not inherited by the renderer.
        env = {'PATH': '/usr/local/bin:/usr/bin:/bin', 'HOME': str(home), 'LANG': 'C.UTF-8',
               'TMPDIR': str(home)}
        try:
            process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                       env=env, start_new_session=True)
            try:
                output, _ = process.communicate(timeout=cfg.conversion_timeout)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid, signal.SIGKILL)
                process.communicate()
                raise ConversionError('EXCEL CONVERSION TIMED OUT')
        except OSError as exc:
            raise ConversionError('LIBREOFFICE COULD NOT START') from exc
        target = directory / (workbook.stem + '.pdf')
        if process.returncode != 0 or not target.is_file():
            raise ConversionError('EXCEL CONVERSION FAILED')
        return target
