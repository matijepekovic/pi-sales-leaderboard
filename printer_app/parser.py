"""Read values, never Office code. Discover the largest unambiguous Sub Status table."""
from __future__ import annotations

import logging
import re
import zipfile
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

log = logging.getLogger(__name__)
MAX_ROWS, MAX_COLUMNS, MAX_CELLS = 50000, 256, 500000


class ParseError(ValueError):
    pass


@dataclass
class Group:
    key: str
    label: str
    headers: list
    rows: list[list]
    error: str = ''


def nonblank(value) -> bool:
    return value is not None and str(value).strip() != ''


def normalized(value) -> str:
    return re.sub(r'[^a-z0-9]', '', str(value).casefold())


def read_sheets(path: Path):
    if path.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            total = 0
            for sheet in book.sheets():
                total += sheet.nrows * sheet.ncols
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS or total > MAX_CELLS:
                    raise ParseError('WORKBOOK EXCEEDS SAFE ROW/COLUMN LIMIT')
                rows = []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            raise ParseError('EXCEL CELL ERROR FOUND')
                        row.append(value)
                    rows.append(row)
                yield sheet.name, rows
        finally:
            book.release_resources()
        return
    with zipfile.ZipFile(path) as archive:
        info = archive.infolist()
        if len(info) > 10000 or sum(i.file_size for i in info) > 128 * 1024 * 1024:
            raise ParseError('WORKBOOK EXCEEDS SAFE EXPANSION LIMIT')
    # VBA, external links, hyperlinks and formulas are never copied into output.
    values = load_workbook(path, read_only=True, data_only=True, keep_links=False, keep_vba=False)
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False, keep_vba=False)
    try:
        total = 0
        for sheet, formula_sheet in zip(values.worksheets, formulas.worksheets):
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ParseError('WORKBOOK EXCEEDS SAFE ROW/COLUMN LIMIT')
            rows = []
            for row, formula_row in zip(sheet.iter_rows(), formula_sheet.iter_rows()):
                total += len(row)
                if total > MAX_CELLS or len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise ParseError('WORKBOOK EXCEEDS SAFE ROW/COLUMN LIMIT')
                vals = []
                for value, formula in zip(row, formula_row):
                    if formula.data_type == 'f' and value.value is None:
                        raise ParseError('FORMULA HAS NO CACHED VALUE; RECALCULATE AND SAVE SOURCE')
                    if value.data_type == 'e':
                        raise ParseError('EXCEL CELL ERROR FOUND')
                    vals.append(value.value)
                rows.append(vals)
            yield sheet.title, rows
    finally:
        values.close()
        formulas.close()


def parse(path: Path) -> list[Group]:
    candidates = []
    try:
        for name, rows in read_sheets(path):
            for r, header in enumerate(rows[:100]):
                columns = [i for i, v in enumerate(header) if normalized(v) == 'substatus']
                if not columns:
                    continue
                if len(columns) != 1:
                    raise ParseError('AMBIGUOUS SUB STATUS COLUMNS')
                used = [i for i, v in enumerate(header) if nonblank(v)]
                left, right = min(used), max(used)
                headers = header[left:right + 1]
                body = []
                for row in rows[r + 1:]:
                    cells = (row + [None] * (right + 1))[left:right + 1]
                    if any(nonblank(v) for v in cells) and cells != headers:
                        body.append(cells)
                candidates.append((len(body), name, r + 1, columns[0] - left, headers, body))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError('EXCEL PARSER FAILED: ' + type(exc).__name__) from exc
    if not candidates:
        raise ParseError('SUB STATUS COLUMN NOT FOUND')
    candidates.sort(key=lambda item: item[0], reverse=True)
    chosen = candidates[0]
    if len(candidates) > 1 and candidates[1][0] == chosen[0]:
        raise ParseError('AMBIGUOUS MAIN TABLE; MULTIPLE TABLES HAVE THE SAME SIZE')
    count, sheet, rownum, col, headers, body = chosen
    if count == 0:
        raise ParseError('NO PRINTABLE ROWS FOUND')
    log.info('Main table: sheet=%r header_row=%s rows=%s columns=%s', sheet, rownum, count, len(headers))
    grouped = OrderedDict()
    for row in body:
        label = str(row[col]).strip() if nonblank(row[col]) else ''
        key = 'value:' + label if label else 'missing:'
        if key not in grouped:
            grouped[key] = Group(key, label, headers, [], 'ROW HAS NO SUB STATUS VALUE' if not label else '')
        grouped[key].rows.append(row)
    log.info('Sub Status groups: %s', len(grouped))
    return list(grouped.values())
