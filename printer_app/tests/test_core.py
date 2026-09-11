from __future__ import annotations

import ast
import base64
import hashlib
import shutil
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from printer_app import converter
from printer_app.print_options import PrintOptions
from dataclasses import replace
from printer_app.config import Config, safe_name
from printer_app.db import Database
from printer_app.deploy import fingerprint, source_files
from printer_app.error_pages import error_page
from printer_app.gmail_client import GmailClient, attachment_parts, sexpr
from printer_app.printer import MissingJob, PrinterError, SubmissionRejected
from printer_app.worker import Engine


def pdf(path: Path, pages=1):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=792, height=612)
    writer.write(path)
    return path


def workbook(path: Path, rows=None):
    book = Workbook()
    sheet = book.active
    sheet.append(['A title above the table'])
    sheet.append([])
    sheet.append(['Name', 'Sub Status', 'Amount'])
    sheet.column_dimensions['A'].width = 30
    sheet.column_dimensions['B'].width = 30
    for row in rows or [['A', 'Arbitrary / QC', 12], ['B', 'Waiting for XYZ', 15], ['C', 'Arbitrary / QC', 20]]:
        sheet.append(row)
    book.save(path)
    return path


class FakePrinter:
    def __init__(self):
        self.jobs, self.printed = {}, []
        self.fail_hold = False
        self.offline = False

    def find(self, token):
        if self.offline:
            raise PrinterError('CUPS CONNECTION FAILED')
        return [dict(value, **{'job-id': key}) for key, value in self.jobs.items() if value['job-name'] == token]

    def hold(self, path, token, options, *, received_pdf=False):
        if self.fail_hold:
            self.fail_hold = False
            raise SubmissionRejected('PRINTER SUBMISSION FAILED')
        jid = len(self.jobs) + 1
        self.jobs[jid] = {'job-name': token, 'job-state': 4, 'path': str(path), 'options': options, 'received_pdf': received_pdf}
        return jid, f'request id is konicaa-{jid} (1 file(s))', ['lp', '-H', 'hold', '--', str(path)]

    def attributes(self, jid):
        if jid not in self.jobs:
            raise MissingJob('CUPS JOB HISTORY IS UNAVAILABLE')
        return self.jobs[jid]

    def release(self, jid):
        if self.jobs[jid]['job-state'] == 4:
            self.printed.append(self.jobs[jid]['path'])
            self.jobs[jid]['job-state'] = 9
        return 'released'

    def cancel_held_duplicate(self, jid):
        self.jobs[jid]['job-state'] = 7


@pytest.fixture
def rig(tmp_path):
    cfg = Config(data_dir=tmp_path / 'private')
    db = Database(cfg.db_path)
    printer = FakePrinter()
    return cfg, db, printer, Engine(cfg, db, printer)


def attach(rig, path: Path, part='1'):
    cfg, db, _, _ = rig
    identity = hashlib.sha256((str(path) + part).encode()).hexdigest()
    mid = db.execute('''INSERT INTO processed_messages
        (identity,account,mailbox,uidvalidity,uid,message_id,subject,sender,created)
        VALUES (?,'test','INBOX','1',?,?,'Test subject','sender@example.test',?)''',
        (identity, identity, identity, time.time()))
    GmailClient(cfg, db)._store(mid, part, path.name, path.read_bytes())
    return db.one('SELECT * FROM attachments WHERE message_id=?', (mid,))


def complete(rig):
    _, db, _, engine = rig
    for _ in range(4):
        for job in db.rows("SELECT * FROM jobs WHERE status IN ('READY','PRINTER ERROR','SUBMITTED')"):
            engine.advance(job)


@pytest.mark.parametrize('pages', [1, 2, 5])
def test_incoming_pdfs_print_directly_even_multipage(rig, tmp_path, pages):
    cfg, db, printer, engine = rig
    source = pdf(tmp_path / 'direct.pdf', pages)
    attachment = attach(rig, source)
    engine.process_attachment(attachment)
    complete(rig)
    job = db.recent()[0]
    assert job['status'] == 'PRINTED' and job['page_count'] == pages
    assert Path(printer.printed[0]).read_bytes() == source.read_bytes()
    assert not job['tabloid']


@pytest.mark.parametrize('filename,payload', [('unknown.doc', b'not supported'), ('broken.xlsx', b'corrupt'), ('broken.pdf', b'corrupt')])
def test_bad_attachments_print_one_error_sheet(rig, tmp_path, filename, payload):
    _, db, printer, engine = rig
    path = tmp_path / filename
    path.write_bytes(payload)
    engine.process_attachment(attach(rig, path))
    complete(rig)
    assert db.recent()[0]['status'] == 'ERROR PRINTED'
    assert converter.page_count(Path(printer.printed[0])) == 1
    assert 'PRINT ERROR' in PdfReader(printer.printed[0]).pages[0].extract_text()


def test_whole_workbook_is_one_job_without_splitting(rig, tmp_path, monkeypatch):
    _, db, printer, engine = rig
    source = workbook(tmp_path / 'groups.xlsx')
    original = source.read_bytes()
    def render(path, directory, cfg):
        assert path.read_bytes() == original
        return pdf(directory / 'report.pdf')
    monkeypatch.setattr(converter, 'convert', render)
    engine.process_attachment(attach(rig, source))
    complete(rig)
    assert len(db.recent()) == len(printer.printed) == 1
    assert db.recent()[0]['substatus'] == ''
    assert db.recent()[0]['group_key'] == 'workbook'
    assert source.read_bytes() == original


@pytest.mark.parametrize('policy,expected', [('one-page', 'ERROR PRINTED'), ('all', 'PRINTED')])
def test_excel_page_policy(rig, tmp_path, monkeypatch, policy, expected):
    cfg, db, printer, _ = rig
    engine = Engine(replace(cfg, print_options=PrintOptions(page_policy=policy)), db, printer)
    monkeypatch.setattr(converter, 'convert', lambda path, directory, cfg: pdf(directory / 'report.pdf', 3))
    engine.process_attachment(attach(rig, workbook(tmp_path / 'long.xlsx')))
    complete((cfg, db, printer, engine))
    assert db.recent()[0]['status'] == expected
    assert db.recent()[0]['page_count'] == 3
    assert converter.page_count(Path(printer.printed[0])) == (1 if policy == 'one-page' else 3)


def test_one_conversion_exception_does_not_abort_other_attachments(rig, tmp_path, monkeypatch):
    _, db, printer, engine = rig
    monkeypatch.setattr(converter, 'convert', lambda *a: (_ for _ in ()).throw(converter.ConversionError('EXCEL CONVERSION FAILED')))
    engine.process_attachment(attach(rig, workbook(tmp_path / 'broken.xlsx')))
    engine.process_attachment(attach(rig, pdf(tmp_path / 'okay.pdf')))
    complete(rig)
    assert {j['status'] for j in db.recent()} == {'PRINTED', 'ERROR PRINTED'}


def test_saved_job_settings_survive_restart_and_later_changes(rig, tmp_path, monkeypatch):
    cfg, db, printer, _ = rig
    first = PrintOptions(paper='tabloid', orientation='landscape', copies=3, sides='two-sided-short-edge')
    engine = Engine(replace(cfg, print_options=first), db, printer)
    engine.process_attachment(attach(rig, pdf(tmp_path / 'original.pdf')))
    later = Engine(replace(cfg, print_options=PrintOptions(copies=10, color='monochrome')), db, printer)
    complete((cfg, db, printer, later))
    assert printer.jobs[1]['options'] == first
    assert Database(cfg.db_path).job_print_settings(db.recent()[0]['id']) == first.snapshot()


def test_legacy_partial_workbook_never_reprints_as_full_report(rig, tmp_path, monkeypatch):
    _, db, _, engine = rig
    attachment = attach(rig, workbook(tmp_path / 'legacy.xlsx'))
    now = time.time()
    db.execute("INSERT INTO jobs(attachment_id,group_key,status,created,updated) VALUES (?,'value:Old','PRINTED',?,?)", (attachment['id'], now, now))
    db.execute("INSERT INTO jobs(attachment_id,group_key,status,created,updated) VALUES (?,'value:Next','PREPARING',?,?)", (attachment['id'], now, now))
    monkeypatch.setattr(converter, 'convert', lambda *a: pytest.fail('Do not print the full workbook again'))
    engine.process_attachment(attachment)
    complete(rig)
    assert len(db.recent()) == 2
    assert {j['status'] for j in db.recent()} == {'PRINTED', 'ERROR PRINTED'}


def test_completed_attachment_restart_does_not_reprint(rig, tmp_path):
    cfg, db, printer, engine = rig
    attachment = attach(rig, pdf(tmp_path / 'direct.pdf'))
    engine.process_attachment(attachment)
    complete(rig)
    Engine(cfg, Database(cfg.db_path), printer).process_attachment(attachment)
    complete(rig)
    assert len(printer.printed) == 1 and len(db.recent()) == 1


def test_crash_after_cups_hold_before_sqlite_receipt(rig, tmp_path):
    cfg, db, printer, engine = rig
    engine.process_attachment(attach(rig, pdf(tmp_path / 'direct.pdf')))
    job = db.recent()[0]
    token = 'printer-app-crash-test'
    db.execute('INSERT INTO print_attempts(job_id,token,created,updated) VALUES (?,?,?,?)', (job['id'], token, time.time(), time.time()))
    printer.hold(Path(job['printable']), token, PrintOptions())
    restarted = Engine(cfg, Database(cfg.db_path), printer)
    restarted.advance(db.job(job['id']))
    restarted.advance(db.job(job['id']))
    assert len(printer.jobs) == 1 and len(printer.printed) == 1
    assert db.job(job['id'])['status'] == 'PRINTED'


def test_receipt_is_committed_before_physical_release(rig, tmp_path):
    _, db, printer, engine = rig
    engine.process_attachment(attach(rig, pdf(tmp_path / 'direct.pdf')))
    original = printer.release
    def release(jid):
        row = db.one('SELECT * FROM print_attempts WHERE cups_id=?', (jid,))
        assert row and row['request_id'] == 'konicaa-1' and row['state'] == 'RELEASING'
        return original(jid)
    printer.release = release
    complete(rig)
    assert len(printer.printed) == 1


def test_printer_rejection_becomes_error_sheet_without_recursion(rig, tmp_path):
    _, db, printer, engine = rig
    engine.process_attachment(attach(rig, pdf(tmp_path / 'direct.pdf')))
    printer.fail_hold = True
    complete(rig)
    assert db.recent()[0]['status'] == 'ERROR PRINTED'
    assert len(printer.printed) == 1 and Path(printer.printed[0]).name == 'error.pdf'


def test_printer_outage_keeps_durable_pending_job(rig, tmp_path):
    _, db, printer, engine = rig
    engine.process_attachment(attach(rig, pdf(tmp_path / 'direct.pdf')))
    printer.offline = True
    complete(rig)
    assert db.recent()[0]['status'] == 'PRINTER ERROR' and not printer.printed
    printer.offline = False
    complete(rig)
    assert len(printer.printed) == 1 and db.recent()[0]['status'] == 'PRINTED'


def test_lost_cups_history_never_blindly_reprints_original(rig, tmp_path):
    _, db, printer, engine = rig
    engine.process_attachment(attach(rig, pdf(tmp_path / 'direct.pdf')))
    engine.advance(db.recent()[0])
    printer.jobs.clear()
    engine.advance(db.recent()[0])
    assert db.recent()[0]['is_error']
    complete(rig)
    assert len(printer.printed) == 2
    assert Path(printer.printed[0]).name == 'direct.pdf' and Path(printer.printed[1]).name == 'error.pdf'


def test_long_error_text_stays_one_page(tmp_path):
    target = error_page(tmp_path / 'error.pdf', dict(id=3, filename='X' * 900, subject='S' * 2000,
        sender='email ' * 300, substatus='Y' * 900), 'CONVERSION ERROR ' * 500, 'America/Los_Angeles')
    assert converter.page_count(target) == 1


def test_filename_cannot_escape_or_be_a_command():
    assert safe_name('../../-bad.pdf') == 'bad.pdf'
    assert '/' not in safe_name('x/../../file.xlsm')
    assert safe_name('...') == 'attachment'


def test_imap_bodystructure_nested_and_filename():
    tree = sexpr(b'(("TEXT" "PLAIN" ("CHARSET" "UTF-8") NIL NIL "7BIT" 12 1 NIL NIL) ("APPLICATION" "PDF" ("NAME" "test.pdf") NIL NIL "BASE64" 200 NIL ("ATTACHMENT" ("FILENAME" "test.pdf"))) "MIXED")')
    parts = list(attachment_parts(tree))
    assert len(parts) == 1 and parts[0]['part'] == '2' and parts[0]['filename'] == 'test.pdf'


def test_gmail_dedup_uid_and_message_id(rig, tmp_path, monkeypatch):
    cfg, db, printer, engine = rig
    from dataclasses import replace
    cfg = replace(cfg, email_user='test@example.test', email_password='not-a-real-secret')
    encoded = base64.b64encode(pdf(tmp_path / 'source.pdf').read_bytes())
    class IMAP:
        uids = [b'1']
        fetches = 0
        def __init__(self, *args, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def login(self, *args): return 'OK', []
        def select(self, mailbox, readonly): assert readonly; return 'OK', []
        def response(self, name): return name, [b'123']
        def uid(self, action, *args):
            if action == 'search': return 'OK', [b' '.join(self.uids)]
            query = args[1]
            if 'HEADER.FIELDS' in query:
                return 'OK', [(b'x', b'Subject: Example\r\nFrom: from@example.test\r\nMessage-ID: <same@test>\r\n\r\n')]
            if 'BODYSTRUCTURE' in query:
                return 'OK', [b'1 (BODYSTRUCTURE ("APPLICATION" "PDF" NIL NIL NIL "BASE64" 100 NIL ("ATTACHMENT" ("FILENAME" "file.pdf"))))']
            type(self).fetches += 1
            return 'OK', [(b'x', encoded)]
    monkeypatch.setattr('printer_app.gmail_client.imaplib.IMAP4_SSL', IMAP)
    gmail = GmailClient(cfg, db)
    assert gmail.poll() == 1
    assert gmail.poll() == 0
    IMAP.uids = [b'1', b'2']
    assert gmail.poll() == 0
    assert IMAP.fetches == 1 and len(db.rows('SELECT * FROM attachments')) == 1
    engine.process_attachment(db.one('SELECT * FROM attachments'))
    complete(rig)
    assert len(printer.printed) == 1


def test_gmail_unconfigured_does_not_crash(rig):
    cfg, db, _, _ = rig
    assert GmailClient(cfg, db).poll() == 0
    assert db.get('gmail_state') == 'NOT CONFIGURED'


def test_xls_support(tmp_path):
    xlwt = pytest.importorskip('xlwt')
    if not shutil.which('libreoffice'):
        pytest.skip('LibreOffice not installed')
    book = xlwt.Workbook()
    sheet = book.add_sheet('Report')
    sheet.write(0, 0, 'Legacy workbook, no required columns')
    path = tmp_path / 'old.xls'
    book.save(str(path))
    original = path.read_bytes()
    result = converter.convert(path, tmp_path / 'result', Config(data_dir=tmp_path))
    assert 'Legacy workbook' in PdfReader(result).pages[0].extract_text()
    assert path.read_bytes() == original


@pytest.mark.parametrize('suffix', ['.xlsx', '.xlsm'])
def test_original_workbook_tabloid_conversion(tmp_path, suffix):
    if not shutil.which('libreoffice'):
        pytest.skip('LibreOffice not installed')
    cfg = Config(data_dir=tmp_path, print_options=PrintOptions(paper='tabloid', orientation='landscape', excel_scale='fit-page'))
    source = workbook(tmp_path / ('original' + suffix))
    before = source.read_bytes()
    result = converter.convert(source, tmp_path / 'result', cfg)
    assert converter.page_count(result) == 1
    page = PdfReader(result).pages[0]
    assert abs(float(page.mediabox.width) - 1224) < 2
    assert abs(float(page.mediabox.height) - 792) < 2
    assert 'A title above the table' in page.extract_text()
    assert 'Arbitrary / QC' in page.extract_text()
    assert 'Waiting for XYZ' in page.extract_text()
    assert source.read_bytes() == before


def test_runtime_secrets_and_data_are_excluded_from_release(tmp_path):
    (tmp_path / '.env').write_text('secret')
    (tmp_path / '.env.example').write_text('EMAIL_USER=')
    (tmp_path / 'module.py').write_text('pass')
    (tmp_path / 'data').mkdir()
    (tmp_path / 'data/secret.db').write_text('secret')
    files = [str(rel) for _, rel in source_files(tmp_path)]
    assert files == ['.env.example', 'module.py']
    first = fingerprint(tmp_path)
    (tmp_path / '.env').write_text('changed secret')
    assert fingerprint(tmp_path) == first


def test_no_stats_imports_or_service_dependencies():
    root = Path(__file__).resolve().parents[1]
    forbidden = {'stats_core', 'database', 'server', 'tableau_scheduler', 'themes', 'source_picker'}
    for path in root.glob('*.py'):
        tree = ast.parse(path.read_text())
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                assert not {a.name.split('.')[0] for a in node.names} & forbidden
            if isinstance(node, ast.ImportFrom) and node.level == 0:
                assert (node.module or '').split('.')[0] not in forbidden
            if isinstance(node, ast.Call):
                assert not any(k.arg == 'shell' and isinstance(k.value, ast.Constant) and k.value.value is True for k in node.keywords)
    for path in (root / 'systemd').glob('*.service'):
        text = path.read_text()
        assert 'User=scoreboard' in text
        assert not any(name in text for name in ('pi-tableau-leaderboard', 'Requires=', 'BindsTo=', 'PartOf='))
