"""Actual XLSX -> existing Office adapter -> notes -> durable queue regressions.

Only the external report download is a fake. Fixtures contain invented job data.
"""
from dataclasses import replace
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from pypdf import PdfReader

from printer_app.report_notes_contract import ReportNotesOptions

REPORT_ID = '00O123456789012'


def workbook_bytes(many=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Jobs'
    sheet.append(['Daily jobs'])
    sheet.merge_cells('A1:C1')
    sheet.append(['Filtered by: Example market'])
    sheet.merge_cells('A2:C2')
    sheet.append(['Customer', 'Job', 'Product'])
    for index in range(60 if many else 3):
        sheet.append([f'Example {index:03d}', str(1000 + index), 'Roof' if index % 2 else 'Bath'])
    # Hidden records are not printed and must not appear in the overlay either.
    sheet.row_dimensions[5].hidden = True
    sheet.append(['Subtotal', 'Count', 59 if many else 2])
    edge = Side(style='thin', color='999999')
    for row in sheet:
        for cell in row:
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.font = Font(name='Liberation Sans', size=10, bold=cell.row <= 3)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 26
    for key, width in (('A', 28), ('B', 18), ('C', 30)):
        sheet.column_dimensions[key].width = width
    sheet.print_area = f'A1:C{sheet.max_row}'
    sheet.print_title_rows = '3:3'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0 if many else 1
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.25, bottom=.25)
    sheet.oddFooter.center.text = 'Test page &P'
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


class Source:
    def __init__(self, many=False):
        self.payload = workbook_bytes(many)
        self.calls = []

    def download_formatted_report(self, report_id):
        self.calls.append(report_id)
        return self.payload


class NotesPersistenceTests(unittest.TestCase):
    def test_legacy_payload_and_notes_roundtrip(self):
        from printer_app.report_printing_contract import ReportPrintingJob, ScheduledReport
        from printer_app.report_printing_repository import ReportPrintingRepository
        from printer_app.print_options import PrintOptions
        from printer_app.db import Database
        notes = ReportNotesOptions(True, ('Job', 'Customer'), 14, 11.5, 3)
        job = ReportPrintingJob('test', 'Morning', 'UTC', 8, 0, (0, 1), PrintOptions(),
                               (ScheduledReport(REPORT_ID, 'Daily', {}, notes),))
        encoded = ReportPrintingRepository.encode(job)
        self.assertEqual(ReportPrintingRepository.decode(encoded), job)
        del encoded['reports'][0]['notes']
        self.assertFalse(ReportPrintingRepository.decode(encoded).reports[0].notes.enabled)
        with tempfile.TemporaryDirectory() as temp:
            repository = ReportPrintingRepository(Database(Path(temp) / 'printer.db'))
            repository.save(job)
            self.assertEqual(repository.list(), [job])


@unittest.skipUnless(shutil.which('libreoffice'), 'LibreOffice not installed')
class NotesOfficeIntegrationTests(unittest.TestCase):
    def setUp(self):
        from printer_app.config import Config
        from printer_app.print_options import PrintOptions
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.cfg = Config(data_dir=self.root, email_enabled=False)
        self.options = PrintOptions(paper='letter', orientation='portrait',
                                    excel_scale='fit-width', margins='narrow',
                                    page_policy='all', color='monochrome')
        self.notes = ReportNotesOptions(True, ('Customer', 'Job'))

    def test_real_conversion_preserves_source_and_hidden_rows(self):
        from printer_app.report_printing_service import prepare_report
        from printer_app.report_notes_pdf import normalize_pdf
        from printer_app.report_table_pdf import read_tables
        payload = workbook_bytes()
        pdf, pages, placements = prepare_report(payload, self.root, self.cfg, self.options, self.notes)
        self.assertEqual((self.root / 'source.xlsx').read_bytes(), payload)
        self.assertEqual(pdf.name, 'report-notes.pdf')
        self.assertEqual(pages, 1)
        self.assertEqual(sum(p.rows for p in placements), 2)
        native = read_tables(normalize_pdf((self.root / 'report.pdf').read_bytes()), self.notes.columns)
        self.assertEqual(native.pages[0].rows, (('Example 000', '1000'), ('Example 002', '1002')))
        text = PdfReader(pdf).pages[0].extract_text()
        self.assertEqual(text.count('Example 000'), 2)
        self.assertEqual(text.count('Example 002'), 2)
        self.assertNotIn('Example 001', text)
        self.assertGreater(placements[0].notes_top, native.pages[0].bottom)

    def test_disabled_preparation_is_original_pdf_without_parsing(self):
        from printer_app.report_printing_service import prepare_report
        with patch('printer_app.report_table_pdf.read_tables', side_effect=AssertionError('disabled parser')):
            pdf, _, placements = prepare_report(workbook_bytes(), self.root, self.cfg, self.options, ReportNotesOptions())
        self.assertEqual(pdf, self.root / 'report.pdf')
        self.assertEqual(placements, ())
        self.assertNotIn('Updates', PdfReader(pdf).pages[0].extract_text())

    def test_long_report_matches_rows_to_real_pages_and_shrinks(self):
        from printer_app.report_printing_service import prepare_report
        from printer_app.report_notes_pdf import normalize_pdf
        from printer_app.report_table_pdf import read_tables
        notes = replace(self.notes, row_spacing_mm=18)
        pdf, pages, placements = prepare_report(workbook_bytes(True), self.root, self.cfg, self.options, notes)
        original = (self.root / 'report.pdf').read_bytes()
        parsed = read_tables(normalize_pdf(original), notes.columns)
        self.assertGreater(pages, 1)
        self.assertEqual(pages, len(PdfReader(BytesIO(original)).pages))
        self.assertEqual(sum(len(p.rows) for p in parsed.pages if p), 59)
        self.assertEqual(sum(p.rows for p in placements), 59)
        self.assertTrue(any(p.scale < 1 for p in placements))
        for placement, source_page, final_page in zip(placements, parsed.pages, PdfReader(pdf).pages):
            self.assertLessEqual(placement.scale, 1)
            self.assertGreater(placement.notes_top, placement.table_bottom)
            self.assertLessEqual(placement.notes_bottom, placement.footer_top + .01)
            printed = final_page.extract_text()
            for name, number in source_page.rows:
                self.assertIn(name, printed)
                self.assertIn(number, printed)
            other = {row[0] for p in parsed.pages if p is not source_page for row in p.rows}
            self.assertTrue(all(name not in printed for name in other))

    def test_scheduler_enqueues_overlay_once_with_existing_print_options(self):
        from printer_app.db import Database
        from printer_app.report_printing_contract import ReportPrintingJob, ScheduledReport
        from printer_app.report_printing_repository import ReportPrintingRepository
        from printer_app.report_printing_service import ReportPrintingService
        db = Database(self.cfg.db_path)
        repository = ReportPrintingRepository(db)
        source = Source()
        service = ReportPrintingService(repository, source, db, self.cfg)
        job = ReportPrintingJob('test', 'Morning', 'UTC', 8, 0, (0,), self.options,
                               (ScheduledReport(REPORT_ID, 'Daily', {}, self.notes),))
        service.run(job, '2026-10-09T08:00:00+00:00')
        service.run(job, '2026-10-09T08:00:00+00:00')
        run = repository.runs()[0]
        self.assertEqual(run['state'], 'QUEUED', run['error'])
        queued = service.queue.activity_job(run['queued_job_id'])
        self.assertEqual(Path(queued['printable']).name, 'report-notes.pdf')
        self.assertIn('Updates', PdfReader(queued['printable']).pages[0].extract_text())
        settings = db.get('job_print_settings:' + str(run['queued_job_id']))
        if isinstance(settings, str):
            settings = json.loads(settings)
        self.assertEqual(settings, self.options.snapshot())
        self.assertEqual(source.calls, [REPORT_ID])
        self.assertEqual(len(service.queue.recent_activity()), 1)

    def test_bad_columns_do_not_silently_queue_unannotated_report(self):
        from printer_app.db import Database
        from printer_app.report_printing_contract import ReportPrintingJob, ScheduledReport
        from printer_app.report_printing_repository import ReportPrintingRepository
        from printer_app.report_printing_service import ReportPrintingService
        db = Database(self.cfg.db_path)
        repository = ReportPrintingRepository(db)
        service = ReportPrintingService(repository, Source(), db, self.cfg)
        job = ReportPrintingJob('test', 'Morning', 'UTC', 8, 0, (0,), self.options,
            (ScheduledReport(REPORT_ID, 'Daily', {}, replace(self.notes, columns=('Missing',))),))
        service.run(job, 'test-occurrence')
        self.assertEqual(repository.runs()[0]['state'], 'ERROR')
        self.assertEqual(service.queue.recent_activity(), [])

    def test_preview_uses_real_preparation_and_cleans_up_without_enqueue(self):
        from printer_app.report_printing_service import preview_report
        source = Source()
        data, columns, placements = preview_report(source, REPORT_ID, self.cfg, self.options)
        self.assertEqual(columns, ('Customer', 'Job', 'Product'))
        self.assertEqual(placements, ())
        data, columns, placements = preview_report(source, REPORT_ID, self.cfg, self.options, self.notes)
        self.assertEqual(sum(p.rows for p in placements), 2)
        self.assertIn('Updates', PdfReader(BytesIO(data)).pages[0].extract_text())
        self.assertFalse(self.cfg.db_path.exists())
        self.assertEqual(list(self.root.glob('report-preview-*')), [])


if __name__ == '__main__':
    unittest.main()
