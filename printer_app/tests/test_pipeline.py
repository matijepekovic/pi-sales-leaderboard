import subprocess
import time
from pathlib import Path
import pytest
from pypdf import PdfReader
from printer_app.contracts import Message
from printer_app.converter import LibreOfficeRenderer, page_count
from printer_app.parser import parse, ParserError
from printer_app.printer import CupsPrinter
from printer_app.repository import Repository
from printer_app.services import PrintService, ControlService
from printer_app.worker import Worker
from .conftest import make_pdf, ingest, workbook, FakeRenderer


@pytest.mark.parametrize('pages', [1, 2, 7])
def test_received_pdfs_print_directly_even_when_multipage(setup, pages):
    cfg, repo, printer, _, service = setup
    path = make_pdf(cfg.data_dir / 'received.pdf', pages)
    ingest(repo, path)
    service.prepare()
    service.dispatch()
    assert len(printer.calls) == 1
    assert printer.calls[0][0] == path
    assert not printer.calls[0][2]
    assert repo.recent()[0]['state'] == 'SUBMITTED'
    service.dispatch()
    assert repo.recent()[0]['state'] == 'PRINTED'
    assert repo.recent()[0]['page_count'] == pages


@pytest.mark.parametrize('filename,body,reason', [('file.txt', b'text', 'UNSUPPORTED'), ('broken.xlsx', b'broken', 'PARSER FAILED'), ('broken.pdf', b'broken', 'PREFLIGHT FAILED')])
def test_bad_inputs_produce_one_physical_error_sheet(setup, filename, body, reason):
    cfg, repo, printer, _, service = setup
    path = cfg.data_dir / filename
    path.write_bytes(body)
    ingest(repo, path)
    service.prepare()
    service.dispatch()
    service.dispatch()
    assert len(printer.calls) == 1
    assert printer.calls[0][0].name == 'ERROR.pdf'
    assert page_count(printer.calls[0][0]) == 1
    job = repo.recent()[0]
    assert job['state'] == 'ERROR PRINTED'
    assert reason in job['error']
    text = PdfReader(printer.calls[0][0]).pages[0].extract_text()
    assert 'PRINT ERROR' in text and 'Source report' in text and filename in text


def test_groups_are_dynamic_and_failure_is_isolated(setup):
    cfg, repo, printer, _, service = setup
    path = workbook(cfg.data_dir / 'report.xlsx', ['First / custom', 'Fail this', 'Third ???'])
    ingest(repo, path)
    service.renderer = FakeRenderer(fail='Fail this')
    service.prepare()
    service.dispatch()
    service.dispatch()
    jobs = {j['substatus']: j for j in repo.recent()}
    assert len(printer.calls) == 3
    assert jobs['First / custom']['state'] == 'PRINTED'
    assert jobs['Third ???']['state'] == 'PRINTED'
    assert jobs['Fail this']['state'] == 'ERROR PRINTED'


def test_rendered_excel_multipage_only_prints_error(setup):
    cfg, repo, printer, _, service = setup
    path = workbook(cfg.data_dir / 'report.xlsx', ['short', 'long', 'another'])
    ingest(repo, path)
    service.renderer = FakeRenderer(pages={'long': 3})
    service.prepare()
    service.dispatch()
    service.dispatch()
    job = next(j for j in repo.recent() if j['substatus'] == 'long')
    assert job['state'] == 'ERROR PRINTED'
    assert job['page_count'] == 3
    assert '3 PAGES' in job['error']
    outputs = repo.outputs(job['id'])
    report = next(o for o in outputs if o['kind'] == 'REPORT')
    assert report['state'] == 'PREVIEW'
    assert not any(str(call[0]) == report['path'] for call in printer.calls)
    assert all(page_count(call[0]) == 1 for call in printer.calls)


def test_header_detection_column_order_and_blank_status(setup):
    cfg, *_ = setup
    groups = parse(workbook(cfg.data_dir / 'input.xlsx', ['New dynamic status', None, 'New dynamic status']))
    assert groups[0].headers == ('Job', 'Sub Status', 'Amount')
    assert len(groups[0].rows) == 2
    assert groups[1].substatus == '(blank)'
    assert groups[0].rows[0][0].value == 'Row 0'


def test_duplicate_messages_remain_deduplicated_after_restart(setup):
    cfg, repo, printer, renderer, service = setup
    path = make_pdf(cfg.data_dir / 'input.pdf')
    assert ingest(repo, path)
    service.prepare()
    service.dispatch()
    repo2 = Repository(cfg.database)
    restarted = PrintService(cfg, repo2, renderer, printer)
    assert not ingest(repo2, path)
    assert not ingest(repo2, path, source='different-uid-same-message')
    restarted.prepare()
    restarted.dispatch()
    assert len(printer.calls) == 1
    # Identical file bytes in a genuinely new message are not globally suppressed.
    assert ingest(repo2, path, source='uid-3', message_id='<new-message>')
    restarted.prepare()
    restarted.dispatch()
    assert len(printer.calls) == 2


def test_recover_accepted_submission_after_worker_crash(setup):
    cfg, repo, printer, _, service = setup
    ingest(repo, make_pdf(cfg.data_dir / 'input.pdf'))
    service.prepare()
    output = repo.outputs()[0]
    _, token = repo.begin_attempt(output['id'])
    printer.tokens[token] = 'konicaa-42'
    service.dispatch()
    service.dispatch()
    assert not printer.calls
    assert repo.recent()[0]['state'] == 'PRINTED'
    assert repo.output(output['id'])['request_id'] == 'konicaa-42'


def test_ambiguous_submission_is_never_blindly_duplicated(setup):
    cfg, repo, printer, _, service = setup
    ingest(repo, make_pdf(cfg.data_dir / 'input.pdf'))
    service.prepare()
    printer.uncertain = True
    service.dispatch()
    service.dispatch()
    service.dispatch()
    assert len(printer.calls) == 1
    assert repo.recent()[0]['state'] == 'PRINTER ERROR'
    assert repo.outputs()[0]['state'] == 'UNCERTAIN'


def test_printer_outage_queues_error_and_retries_after_recovery(setup):
    cfg, repo, printer, _, service = setup
    ingest(repo, make_pdf(cfg.data_dir / 'input.pdf'))
    service.prepare()
    printer.reject = True
    service.dispatch()
    service.dispatch()
    assert repo.recent()[0]['state'] == 'PRINTER ERROR'
    output = next(o for o in repo.outputs() if o['kind'] == 'ERROR')
    assert output['state'] == 'RETRY'
    printer.reject = False
    repo.set_output(output['id'], 'RETRY', delay=-1)
    service.dispatch()
    service.dispatch()
    assert repo.recent()[0]['state'] == 'ERROR PRINTED'
    assert sum(call[0].name != 'ERROR.pdf' for call in printer.calls) == 1


def test_cups_acceptance_is_not_completion(setup):
    cfg, repo, printer, _, service = setup
    ingest(repo, make_pdf(cfg.data_dir / 'input.pdf'))
    service.prepare()
    printer.outcome = 'WAITING'
    service.dispatch()
    for _ in range(4):
        service.dispatch()
    assert len(printer.calls) == 1
    assert repo.recent()[0]['state'] == 'SUBMITTED'
    assert repo.get('last_successful_print') is None


def test_cups_unknown_history_is_not_false_success(setup):
    cfg, repo, printer, _, service = setup
    ingest(repo, make_pdf(cfg.data_dir / 'input.pdf'))
    service.prepare()
    service.dispatch()
    printer.outcome = 'UNKNOWN'
    service.dispatch()
    assert repo.recent()[0]['state'] == 'PRINTER ERROR'
    assert repo.get('last_successful_print') is None


def test_lp_uses_argument_array_and_preserves_queue_defaults(tmp_path):
    calls = []
    def runner(args, **kwargs):
        calls.append((args, kwargs))
        return subprocess.CompletedProcess(args, 0, stdout='request id is konicaa-7 (1 file(s))', stderr='')
    adapter = CupsPrinter('konicaa', connection_factory=lambda: None, runner=runner)
    file = tmp_path / 'test;touch BAD.pdf'
    result = adapter.submit(file, 'printer-app-test', False)
    assert result.request_id == 'konicaa-7'
    args, kwargs = calls[0]
    assert args == ['lp', '-d', 'konicaa', '-t', 'printer-app-test', str(file)]
    assert not kwargs.get('shell', False)
    assert 'EMAIL_APP_PASSWORD' not in kwargs['env']


def test_gmail_outage_does_not_crash_worker_and_pause_persists(setup):
    cfg, repo, printer, _, service = setup
    class BrokenMail:
        def messages(self, seen):
            raise ConnectionError('unreachable')
    worker = Worker(cfg, repo, BrokenMail(), service)
    worker.tick()
    assert repo.get('monitor')['state'] == 'error'
    assert repo.get('printer')['known']
    ControlService(cfg, repo).command('pause')
    worker2 = Worker(cfg, Repository(cfg.database), BrokenMail(), service)
    worker2.tick()
    assert repo.get('paused') is True
    ControlService(cfg, repo).command('test')
    worker2.tick()
    assert len(printer.calls) == 1


def test_test_command_is_idempotent(setup):
    _, repo, printer, _, service = setup
    service.test_print('fixed-command-id')
    service.test_print('fixed-command-id')
    service.dispatch()
    service.dispatch()
    assert len(printer.calls) == 1
    assert len(repo.recent()) == 1


@pytest.mark.parametrize('mode,expected', [('missing','SUB STATUS COLUMN NOT FOUND'), ('empty','NO PRINTABLE ROWS FOUND'), ('formula','FORMULA HAS NO CACHED VALUE')])
def test_parser_failures_are_explicit(setup, mode, expected):
    from openpyxl import Workbook
    cfg, *_ = setup
    book = Workbook()
    book.active.append(['Job', 'Status' if mode == 'missing' else 'Sub Status'])
    if mode != 'empty':
        book.active.append(['=1+1' if mode == 'formula' else 'A', 'anything'])
    path = cfg.data_dir / 'bad.xlsx'
    book.save(path)
    book.close()
    with pytest.raises(ParserError, match=expected):
        parse(path)


def test_legacy_xls_reads_without_office_execution(setup):
    import xlwt
    cfg, *_ = setup
    book = xlwt.Workbook()
    sheet = book.add_sheet('Report')
    for row, values in enumerate([['Job', 'Sub Status'], ['A', 'Legacy custom'], ['B', 'Another custom']]):
        for column, value in enumerate(values):
            sheet.write(row, column, value)
    path = cfg.data_dir / 'legacy.xls'
    book.save(str(path))
    assert [g.substatus for g in parse(path)] == ['Legacy custom', 'Another custom']


def test_xlsm_is_read_as_values_only(setup):
    cfg, *_ = setup
    path = workbook(cfg.data_dir / 'macro-enabled.xlsm')
    assert len(parse(path)) == 2


def test_real_libreoffice_tabloid_output_and_values_only_safety(setup):
    from openpyxl import load_workbook
    from printer_app.contracts import Cell, ReportGroup
    cfg, *_ = setup
    group = ReportGroup('test', 'Novel status', ('Job', 'Sub Status', 'Amount'), ((Cell('=HYPERLINK("https://invalid.test","x")'), Cell('Novel status'), Cell(123.45, '$#,##0.00')),), 'Main')
    workbook_path, pdf = LibreOfficeRenderer().render(group, cfg.data_dir / 'real-output')
    assert page_count(pdf) == 1
    page = PdfReader(pdf).pages[0]
    assert abs(float(page.mediabox.width) - 17 * 72) < 2
    assert abs(float(page.mediabox.height) - 11 * 72) < 2
    normalized = load_workbook(workbook_path, data_only=False)
    assert normalized.active['A2'].data_type == 's'
    assert normalized.active['C2'].number_format == '$#,##0.00'
    assert not normalized._external_links
    normalized.close()


def test_real_long_workbook_prints_only_one_page_error(setup):
    cfg, repo, printer, _, service = setup
    path = workbook(cfg.data_dir / 'long.xlsx', ['Lots of rows'] * 250)
    ingest(repo, path)
    service.renderer = LibreOfficeRenderer()
    service.prepare()
    service.dispatch()
    service.dispatch()
    assert repo.recent()[0]['page_count'] > 1
    assert repo.recent()[0]['state'] == 'ERROR PRINTED'
    assert len(printer.calls) == 1
    assert printer.calls[0][0].name == 'ERROR.pdf'
    assert page_count(printer.calls[0][0]) == 1
